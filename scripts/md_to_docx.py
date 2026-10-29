# -*- coding: utf-8 -*-
"""Markdown → DOCX 校对预览器（配置驱动，适用于任意书）。

版式与「页码/分页」的差异不再写死在代码里：
  * 分页与页码解析 → `scripts/lib/bookfmt.py` + `config/book_profiles.json`
  * 字体/字号/页边距/缩进/分割线 → 同一配置文件的 `docx` 段

渲染规则：
  `---` / `\\f`             → 真实 Word 分页符（页边界由 profile 判定）
  `· 12 ·` / `<p align>`    → 页脚页码：居中、蓝色、小字
  `> 脚注` / `① 脚注`        → 页脚注块：细分割线 + 悬挂缩进小字
  `# 章` / `## 节`           → 黑体标题
  `- 条目 ...... 12`         → 目录条目：右对齐点线引导
  `**粗**` / `*斜*` / `` `码` `` → 行内格式；行内圈号自动上标

用法:
  python scripts/md_to_docx.py --book 某书
  python scripts/md_to_docx.py --book 某书 --profile formfeed
  python scripts/md_to_docx.py --md a.md --out preview.docx
  python scripts/md_to_docx.py --book 某书 --from 1 --to 20     # 片段
"""
import re
import sys
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Emu, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
from lib import bookfmt as F

LIB = ROOT / 'library'
OUT_DIR = ROOT / 'outputs' / 'docx'

CIRCLED = '①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'
CIRCLED_RE = re.compile(f'([{CIRCLED}])')
INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)')
TOC_DOT_RE = re.compile(r'^(.*?)\s*\.{3,}\s*(\d{1,4})\s*$')
IMG_RE = re.compile(r'^!\[[^\]]*\]\((.*?)\)\s*$')
HTML_CENTER_RE = re.compile(r'^<p\s+align="center"\s*>(.*?)</p>$', re.I | re.S)


def _hex(c, default='000000'):
    try:
        return RGBColor.from_string((c or default).lstrip('#'))
    except Exception:
        return RGBColor.from_string(default)


class Renderer:
    """按配置把「页列表」渲染成书籍版式的 DOCX。"""

    def __init__(self, style):
        self.st = style
        f = style.get('fonts', {})
        self.cn, self.cn_head = f.get('body_cn', '宋体'), f.get('head_cn', '黑体')
        self.latin, self.mono = f.get('latin', 'Times New Roman'), f.get('mono', 'Consolas')
        s = style.get('sizes_pt', {})
        self.body_pt = s.get('body', 10.5)
        self.head1_pt = s.get('head1', 18)
        self.head2_pt = s.get('head2', 13.5)
        self.head3_pt = s.get('head3', 12)
        self.note_pt = s.get('note', 9)
        self.pageno_pt = s.get('pageno', 9)
        self.toc_pt = s.get('toc', 10)
        c = style.get('colors', {})
        self.pageno_color = _hex(c.get('pageno'), '2E74B5')
        self.note_color = _hex(c.get('note'), '333333')
        b = style.get('body', {})
        self.indent_chars = b.get('indent_chars', 2)
        self.line = b.get('line_spacing', 1.5)
        nb = style.get('note_block', {})
        self.rule_ratio = nb.get('rule_ratio', 0.62)
        self.note_hang = Cm(nb.get('hanging_indent_cm', 0.74))
        self.note_line = nb.get('line_spacing', 1.15)
        tl = style.get('toc_list', {})
        self.dot_leader = tl.get('dot_leader', True)
        self.toc_indent = Cm(tl.get('indent_cm', 0.4))
        im = style.get('images', {})
        self.img_max_px = int(im.get('max_px', 1400))
        self.img_quality = int(im.get('jpeg_quality', 82))
        self.img_width = Cm(im.get('width_cm', 12))
        self._img_cache = {}
        pats = style.get('markdown', {}).get('footnote_line', ['^>\\s?', '^\\s*[①-⑳]'])
        self.note_patterns = [re.compile(p) for p in pats]
        # OCR 文本的标题启发式（章节标题常被压成普通段落）
        self.head_rules = [
            (re.compile(r['pattern']), int(r['level']), int(r.get('max_chars', 40)))
            for r in style.get('markdown', {}).get('heading_heuristic', [])
            if r.get('pattern')
        ]
        self.stats = dict(pages=0, notes=0, pagenos=0, headings=0, paras=0,
                          tables=0, images=0)

    # ---------- 基础 ----------
    def font(self, run, size=None, bold=None, italic=None, color=None, cn=None, latin=None):
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color
        latin = latin or self.latin
        run.font.name = latin
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), latin)
        rfonts.set(qn('w:hAnsi'), latin)
        rfonts.set(qn('w:eastAsia'), cn or self.cn)

    @staticmethod
    def para_format(p, first_indent=None, before=0, after=0, line=1.5,
                    align=None, keep_next=False):
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        if first_indent is not None:
            pf.first_line_indent = first_indent
        if align is not None:
            p.alignment = align
        pf.keep_with_next = keep_next

    @staticmethod
    def border(p, color='999999', size=6):
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), str(size))
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), color)
        pbdr.append(bottom)
        ppr.append(pbdr)

    @staticmethod
    def page_break(doc):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.page_break_before = True

    # ---------- 行内 ----------
    def inline(self, p, text, size=None, color=None, bold=False, cn=None):
        size = size or self.body_pt
        for chunk in INLINE_RE.split(text):
            if not chunk:
                continue
            b, i, code = bold, False, False
            if chunk.startswith('**') and chunk.endswith('**') and len(chunk) > 4:
                chunk, b = chunk[2:-2], True
            elif chunk.startswith('`') and chunk.endswith('`') and len(chunk) > 2:
                chunk, code = chunk[1:-1], True
            elif chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
                chunk, i = chunk[1:-1], True

            pos = 0
            for m in CIRCLED_RE.finditer(chunk):
                if m.start() > 0 and chunk[m.start() - 1] not in ' \t':
                    if m.start() > pos:
                        r = p.add_run(chunk[pos:m.start()])
                        self.font(r, size, b, i, color, cn=cn)
                    r = p.add_run(m.group(1))
                    self.font(r, size, b, i, color, cn=cn)
                    r.font.superscript = True
                    pos = m.end()
            if pos < len(chunk):
                r = p.add_run(chunk[pos:])
                self.font(r, size, b, i, color,
                          cn=(self.mono if code else cn),
                          latin=(self.mono if code else None))

    def prep_image(self, path):
        """大图先压成 JPEG（否则单本 DOCX 动辄 40–50MB，不便分发/预览）。"""
        key = str(path)
        if key in self._img_cache:
            return self._img_cache[key]
        try:
            from PIL import Image
            import hashlib
            need = (max(Image.open(path).size) > self.img_max_px
                    or path.stat().st_size > 800_000)
            if not need:
                self._img_cache[key] = path
                return path
            im = Image.open(path).convert('RGB')
            im.thumbnail((self.img_max_px, self.img_max_px), Image.LANCZOS)
            tag = hashlib.sha1((key + str(self.img_max_px)).encode()).hexdigest()[:10]
            out = ROOT / 'tmp' / '.staging' / f'img_{tag}.jpg'
            out.parent.mkdir(parents=True, exist_ok=True)
            im.save(out, 'JPEG', quality=self.img_quality, optimize=True)
            # 压完反而更大（原图本就是高质量小 JPEG）→ 用原图
            if out.stat().st_size >= path.stat().st_size:
                out.unlink(missing_ok=True)
                out = path
            self._img_cache[key] = out
            return out
        except Exception:
            self._img_cache[key] = path
            return path

    def heading_level(self, text):
        """把「第X章 / 第X节 / 前言…」这类短段落识别成标题（OCR 文本需要）。"""
        t = (text or '').strip()
        if not t or len(t) > 45:
            return 0
        for pat, lvl, max_chars in self.head_rules:
            if len(t) <= max_chars and pat.match(t):
                return lvl
        return 0

    # ---------- 块 ----------
    def blocks(self, text, allow_headings=True):
        lines = text.splitlines()
        i, total = 0, len(lines)
        while i < total:
            s = lines[i].rstrip()
            if not s.strip():
                i += 1
                continue

            m = re.match(r'^(#{1,6})\s+(.*)$', s)
            if m:
                yield ('heading', (len(m.group(1)), m.group(2).strip()))
                i += 1
                continue

            m = HTML_CENTER_RE.match(s.strip())
            if m:
                yield ('center', re.sub(r'<[^>]+>', '', m.group(1)).strip())
                i += 1
                continue

            m = IMG_RE.match(s.strip())
            if m:
                yield ('image', m.group(1))
                i += 1
                continue

            if s.lstrip().startswith('|'):
                rows = []
                while i < total and lines[i].lstrip().startswith('|'):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                    i += 1
                rows = [r for r in rows
                        if not all(re.fullmatch(r':?-{2,}:?', c or '') for c in r)]
                if rows:
                    yield ('table', rows)
                continue

            if any(p.match(s) for p in self.note_patterns):
                buf = []
                while i < total:
                    cur = lines[i]
                    if any(p.match(cur) for p in self.note_patterns):
                        buf.append(re.sub(r'^\s*>\s?', '', cur).strip())
                        i += 1
                    elif not cur.strip():
                        if i + 1 < total and any(p.match(lines[i + 1]) for p in self.note_patterns):
                            i += 1
                        else:
                            break
                    else:
                        break
                for b in buf:
                    if b:
                        yield ('note', b)
                continue

            m = re.match(r'^\s*[-*+]\s+(.*)$', s)
            if m:
                yield ('list', m.group(1).strip())
                i += 1
                continue

            # 单行即为章节标题（OCR 文本里标题与正文常无空行分隔）
            lvl = self.heading_level(s) if allow_headings else 0
            if lvl:
                yield ('hheading', (lvl, s.strip()))
                i += 1
                continue

            buf = [s.strip()]
            i += 1
            while i < total:
                nxt = lines[i]
                if (not nxt.strip() or nxt.lstrip().startswith(('#', '>', '|', '-', '!['))
                        or any(p.match(nxt) for p in self.note_patterns)
                        or re.fullmatch(r'\s*-{3,}\s*', nxt)
                        or (allow_headings and self.heading_level(nxt))):  # 下一行是章节标题 → 截断本段
                    break
                buf.append(nxt.strip())
                i += 1
            yield ('para', ''.join(buf))

    # ---------- 渲染 ----------
    def render(self, pages, out_path, title=None):
        doc = Document()
        st = self.st
        page = st.get('page', {})
        mg = st.get('margins_cm', {})
        sec = doc.sections[0]
        sec.page_width = Cm(page.get('width_cm', 18.5))
        sec.page_height = Cm(page.get('height_cm', 26.0))
        sec.top_margin = Cm(mg.get('top', 2.4))
        sec.bottom_margin = Cm(mg.get('bottom', 2.4))
        sec.left_margin = Cm(mg.get('left', 2.2))
        sec.right_margin = Cm(mg.get('right', 2.0))

        normal = doc.styles['Normal']
        normal.font.name = self.latin
        normal.font.size = Pt(self.body_pt)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), self.cn)

        story_w = sec.page_width - sec.left_margin - sec.right_margin
        first_indent = Pt(self.body_pt * self.indent_chars)

        # 扉页：md 的 `# 书名` 不属于任何一页，单独排在最前
        if title:
            cover = st.get('cover', {})
            p = doc.add_paragraph()
            self.para_format(p, before=0, after=0, line=1.5,
                             align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(cover.get('space_before_pt', 180))
            self.inline(p, title, size=cover.get('size_pt', 26), bold=True,
                        cn=self.cn_head)
            self.page_break(doc)
            self.stats['cover'] = 1

        for pi, item in enumerate(pages):
            section = None
            if isinstance(item, (tuple, list)):
                if len(item) >= 3:
                    label, text, section = item[0], item[1], item[2]
                else:
                    label, text = item
            else:
                label, text = item.get('label'), item.get('text')
                section = item.get('section_type')
            note_started = False

            for kind, payload in self.blocks(text or '', allow_headings=(section != 'toc')):
                if kind == 'heading':
                    lvl, title = payload
                    p = doc.add_paragraph()
                    size = {1: self.head1_pt, 2: self.head2_pt,
                            3: self.head3_pt}.get(lvl, 11)
                    if lvl == 1:
                        self.para_format(p, before=6, after=16,
                                         align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
                    else:
                        self.para_format(p, before=12, after=6, keep_next=True)
                    self.inline(p, title, size=size, bold=True, cn=self.cn_head)
                    self.stats['headings'] += 1

                elif kind == 'hheading':
                    lvl, title = payload
                    p = doc.add_paragraph()
                    size = {1: self.head1_pt, 2: self.head2_pt}.get(lvl, self.head3_pt)
                    if lvl == 1:
                        self.para_format(p, before=6, after=16,
                                         align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
                    else:
                        self.para_format(p, before=12, after=6, keep_next=True)
                    self.inline(p, title, size=size, bold=True, cn=self.cn_head)
                    self.stats['headings'] += 1

                elif kind == 'para':
                    lvl = self.heading_level(payload)
                    if lvl:
                        p = doc.add_paragraph()
                        size = {1: self.head1_pt, 2: self.head2_pt}.get(lvl, self.head3_pt)
                        if lvl == 1:
                            self.para_format(p, before=6, after=16,
                                             align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
                        else:
                            self.para_format(p, before=12, after=6, keep_next=True)
                        self.inline(p, payload, size=size, bold=True, cn=self.cn_head)
                        self.stats['headings'] += 1
                    else:
                        p = doc.add_paragraph()
                        self.para_format(p, first_indent=first_indent, line=self.line)
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        self.inline(p, payload)
                        self.stats['paras'] += 1

                elif kind == 'center':
                    p = doc.add_paragraph()
                    self.para_format(p, before=4, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
                    self.inline(p, payload)

                elif kind == 'list':
                    p = doc.add_paragraph()
                    m = TOC_DOT_RE.match(payload)
                    if m and self.dot_leader:
                        self.para_format(p, first_indent=Cm(0), after=2, line=1.35)
                        p.paragraph_format.left_indent = self.toc_indent
                        p.paragraph_format.tab_stops.add_tab_stop(
                            Emu(int(story_w - self.toc_indent)),
                            WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                        self.inline(p, m.group(1).strip(), size=self.toc_pt)
                        r = p.add_run('\t' + m.group(2))
                        self.font(r, self.toc_pt)
                    else:
                        self.para_format(p, first_indent=Cm(0), after=2)
                        p.paragraph_format.left_indent = Cm(0.74)
                        self.inline(p, '· ' + payload)

                elif kind == 'note':
                    if not note_started:
                        note_started = True
                        rule = doc.add_paragraph()
                        self.para_format(rule, before=10, after=2, line=1.0)
                        rule.paragraph_format.right_indent = Emu(int(story_w * self.rule_ratio))
                        self.border(rule)
                    p = doc.add_paragraph()
                    self.para_format(p, first_indent=Cm(0), after=1, line=self.note_line)
                    p.paragraph_format.left_indent = self.note_hang
                    p.paragraph_format.first_line_indent = -self.note_hang
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    self.inline(p, payload, size=self.note_pt, color=self.note_color)
                    self.stats['notes'] += 1

                elif kind == 'table':
                    rows = payload
                    cols = max(len(r) for r in rows)
                    t = doc.add_table(rows=len(rows), cols=cols)
                    t.style = 'Table Grid'
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for ri, row in enumerate(rows):
                        for ci in range(cols):
                            cell = t.cell(ri, ci)
                            cell.text = ''
                            cp = cell.paragraphs[0]
                            self.para_format(cp, first_indent=Cm(0), after=0, line=1.15)
                            self.inline(cp, row[ci] if ci < len(row) else '',
                                        size=9.5, bold=(ri == 0))
                    self.stats['tables'] += 1
                    doc.add_paragraph()

                elif kind == 'image':
                    path = (LIB / payload).resolve()
                    if path.exists():
                        p = doc.add_paragraph()
                        self.para_format(p, first_indent=Cm(0), after=6,
                                         align=WD_ALIGN_PARAGRAPH.CENTER)
                        p.add_run().add_picture(str(self.prep_image(path)),
                                                width=self.img_width)
                        self.stats['images'] += 1

            if label:
                p = doc.add_paragraph()
                self.para_format(p, before=8, after=0, line=1.0,
                                 align=WD_ALIGN_PARAGRAPH.CENTER)
                self.inline(p, label, size=self.pageno_pt, color=self.pageno_color)
                self.stats['pagenos'] += 1

            self.stats['pages'] += 1
            if pi < len(pages) - 1:
                self.page_break(doc)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return self.stats


def md_path_of(book):
    p = LIB / f'{book}.md'
    if p.exists():
        return p
    hits = list(LIB.glob(f'*{book}*.md'))
    if not hits:
        sys.exit(f'找不到 md：{book}')
    return hits[0]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--book')
    ap.add_argument('--md')
    ap.add_argument('--out')
    ap.add_argument('--profile', help='版式名；默认 auto 自动识别')
    ap.add_argument('--config', help='配置文件路径（默认 config/book_profiles.json）')
    ap.add_argument('--from', dest='start', type=int)
    ap.add_argument('--to', dest='end', type=int)
    args = ap.parse_args()

    if not args.book and not args.md:
        ap.error('需要 --book 或 --md')

    src = Path(args.md) if args.md else md_path_of(args.book)
    name = args.book or src.stem

    cfg = F.load_config(args.config)
    md_text = src.read_text(encoding='utf-8')
    pname, prof = F.get_profile(args.profile, md_text, cfg)
    m_title = re.match(r'\s*#\s+([^\n]+)', md_text)
    title = m_title.group(1).strip() if (m_title and not args.start) else None
    pages = F.pages(md_text, prof)

    if args.start or args.end:
        s = (args.start or 1) - 1
        e = args.end or len(pages)
        pages = pages[s:e]

    out = Path(args.out) if args.out else OUT_DIR / f'{name}.docx'
    r = Renderer(F.docx_style(cfg))
    print(f'[DOCX] {src.name} → {out}  （profile={pname}）')
    st = r.render([(p['label'], p['text'], p.get('section_type')) for p in pages],
                  out, title=title)
    cover = '含扉页 | ' if st.get('cover') else ''
    print(f'[DOCX] {cover}页 {st["pages"]} | 页码 {st["pagenos"]} | 脚注 {st["notes"]} | '
          f'标题 {st["headings"]} | 段落 {st["paras"]} | 表格 {st["tables"]} | 图 {st["images"]}')
    print(f'[DOCX] 体积 {out.stat().st_size/1024:.0f} KB')


if __name__ == '__main__':
    main()
